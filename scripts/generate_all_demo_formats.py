import io
import json
import csv
from pathlib import Path

demo_dir = Path(__file__).parent.parent / "data" / "demo_files"
demo_dir.mkdir(parents=True, exist_ok=True)

# 1. DOCX: Commercial Agreement
try:
    import docx
    doc = docx.Document()
    doc.add_heading("MASTER SERVICE AGREEMENT (MSA-2026)", level=1)
    doc.add_paragraph("Paragraph 1: This Agreement is effective as of January 1, 2026 by and between Provider and Client.")
    doc.add_paragraph("Paragraph 2: Section 3.1 Termination Notice: Either party may terminate this Agreement upon providing 30 days prior written notice.")
    doc.add_paragraph("Paragraph 3: Section 4.2 Payment Terms: All undisputed invoices shall be settled within 45 calendar days of receipt.")
    doc.add_paragraph("Paragraph 4: Section 5.1 Service Credits: Failure to achieve 99.9% uptime triggers a fixed service penalty of $15,000 per billing cycle.")
    docx_file = demo_dir / "demo_contract.docx"
    doc.save(str(docx_file))
    print(f"Created: {docx_file.name}")
except Exception as e:
    print(f"DOCX error: {e}")

# 2. PPTX: Executive Presentation
try:
    import pptx
    prs = pptx.Presentation()
    # Slide 1: Title
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Cygnix Executive Strategy 2026"
    s1.shapes.placeholders[1].text = "Trustworthy AI Verification Infrastructure"

    # Slide 2: Milestones
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Key Operational Milestones"
    tf = s2.shapes.placeholders[1].text_frame
    tf.text = "Slide 2: Regulatory annual filings are strictly completed within 30 days of fiscal close."
    p2 = tf.add_paragraph()
    p2.text = "Slide 2: Target annual turnover is projected at $10 million for Enterprise accounts."
    
    # Slide 3: SLA & Penalties
    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    s3.shapes.title.text = "Risk & Penalty Assessment"
    tf3 = s3.shapes.placeholders[1].text_frame
    tf3.text = "Slide 3: The statutory base penalty for late filing violations is $15,000 per violation."
    
    pptx_file = demo_dir / "demo_presentation.pptx"
    prs.save(str(pptx_file))
    print(f"Created: {pptx_file.name}")
except Exception as e:
    print(f"PPTX error: {e}")

# 3. XLSX: Financials
try:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue_Model"
    ws.append(["Category", "Metric", "Statutory Value", "Status"])
    ws.append(["Compliance", "Filing Window", "30 days", "Mandatory"])
    ws.append(["Penalties", "Base Late Fee", "$15,000", "Enforced"])
    ws.append(["Turnover", "Small Enterprise Cap", "$2.5 million", "Exempt"])
    xlsx_file = demo_dir / "demo_financials.xlsx"
    wb.save(str(xlsx_file))
    print(f"Created: {xlsx_file.name}")
except Exception as e:
    print(f"XLSX error: {e}")

# 4. CSV: Records
csv_file = demo_dir / "demo_compliance.csv"
with open(csv_file, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(["Rule_ID", "Requirement_Description", "Timeline", "Statutory_Fine"])
    writer.writerow(["R001", "Corporate annual disclosure submission", "30 days", "$15,000"])
    writer.writerow(["R002", "Quarterly audit exemption for small enterprise", "Under $2.5M", "$0"])
print(f"Created: {csv_file.name}")

# 5. Markdown: Notes
md_file = demo_dir / "demo_notes.md"
md_content = """# Corporate Governance Manual 2026

## Filing Protocols
All mandatory financial reports must be submitted within 30 days of the fiscal year close. Submissions are delivered through the secure electronic portal.

## Penalties and Fines
Any late filing carries a statutory civil penalty of $15,000 per violation.
"""
with open(md_file, "w", encoding="utf-8") as f:
    f.write(md_content)
print(f"Created: {md_file.name}")

# 6. JSON: Config
json_file = demo_dir / "demo_data.json"
json_content = {
    "regulation": "CRD-2026",
    "filing_rules": {
        "mandatory_deadline_days": 30,
        "late_submission_penalty": 15000,
        "small_enterprise_turnover_threshold": 2500000
    }
}
with open(json_file, "w", encoding="utf-8") as f:
    json.dump(json_content, f, indent=2)
print(f"Created: {json_file.name}")

# 7. HTML: Portal
html_file = demo_dir / "demo_portal.html"
html_content = """<!DOCTYPE html>
<html>
<head><title>Compliance Portal</title></head>
<body>
    <header><h1>Regulatory Directive Bulletin</h1></header>
    <article id="filing-guidance">
        <h2>Article 1: Timelines</h2>
        <p>Annual corporate compliance disclosures must be submitted within 30 days of fiscal year end.</p>
        <p>The statutory fine for late submissions is $15,000 per violation.</p>
    </article>
</body>
</html>"""
with open(html_file, "w", encoding="utf-8") as f:
    f.write(html_content)
print(f"Created: {html_file.name}")

print("\nAll demo format files created in data/demo_files/!")
