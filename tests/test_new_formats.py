import io
import json
import csv
from pathlib import Path
import pytest

from backend.services.document_parser import (
    parse_document,
    parse_docx,
    parse_pptx,
    parse_xlsx,
    parse_csv,
    parse_markdown,
    parse_json,
    parse_html
)
from backend.services.verifier import VerificationPipeline
from backend.schemas import VerdictType


def test_docx_extraction():
    import docx
    doc = docx.Document()
    doc.add_paragraph("First contractual clause requires 30-day notice.")
    doc.add_paragraph("Second clause specifies payment in USD.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    meta, chunks = parse_docx(docx_bytes, "agreement.docx")
    assert meta.filename == "agreement.docx"
    assert meta.file_type == "docx"
    assert len(chunks) == 2
    assert chunks[0].location_type == "paragraph"
    assert chunks[0].location == "Paragraph 1"
    assert "30-day notice" in chunks[0].text
    assert chunks[1].location == "Paragraph 2"


def test_pptx_extraction():
    import pptx
    prs = pptx.Presentation()
    # Blank slide layout (index 6)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(0, 0, 100, 100)
    tf = txBox.text_frame
    tf.text = "Q3 Revenue reached $10 million."

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    meta, chunks = parse_pptx(pptx_bytes, "quarterly_deck.pptx")
    assert meta.filename == "quarterly_deck.pptx"
    assert meta.file_type == "pptx"
    assert len(chunks) >= 1
    assert chunks[0].location_type == "slide"
    assert chunks[0].location == "Slide 1"
    assert "$10 million" in chunks[0].text


def test_xlsx_extraction():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws.append(["Metric", "Target", "Actual"])
    ws.append(["Annual Turnover", "$5,000,000", "$6,200,000"])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    meta, chunks = parse_xlsx(xlsx_bytes, "metrics.xlsx")
    assert meta.filename == "metrics.xlsx"
    assert meta.file_type == "xlsx"
    assert len(chunks) >= 2
    assert chunks[1].location_type == "cell"
    assert "Financials!Row 2" in chunks[1].location
    assert "$6,200,000" in chunks[1].text


def test_csv_extraction():
    csv_text = "Employee,Department,Security Clearance\nAlice Smith,Engineering,Top Secret\nBob Jones,Sales,Public\n"
    csv_bytes = csv_text.encode("utf-8")

    meta, chunks = parse_csv(csv_bytes, "staff.csv")
    assert meta.filename == "staff.csv"
    assert meta.file_type == "csv"
    assert len(chunks) == 3
    assert chunks[1].location_type == "row"
    assert chunks[1].location == "Row 2"
    assert "Alice Smith" in chunks[1].text
    assert "Top Secret" in chunks[1].text


def test_markdown_extraction():
    md_text = """# Executive Summary
The corporate mandate requires strict carbon neutrality by 2030.

## Compliance Thresholds
Entities with over 500 staff must publish annual sustainability audits.
"""
    md_bytes = md_text.encode("utf-8")

    meta, chunks = parse_markdown(md_bytes, "sustainability.md")
    assert meta.filename == "sustainability.md"
    assert meta.file_type == "md"
    assert len(chunks) == 2
    assert chunks[0].location_type == "heading"
    assert "Executive Summary" in chunks[0].location
    assert "2030" in chunks[0].text
    assert "Compliance Thresholds" in chunks[1].location


def test_json_extraction():
    json_data = {
        "company": {
            "name": "Acme Corp",
            "statutory_filing_deadline_days": 30,
            "late_penalty": 15000
        }
    }
    json_bytes = json.dumps(json_data).encode("utf-8")

    meta, chunks = parse_json(json_bytes, "config.json")
    assert meta.filename == "config.json"
    assert meta.file_type == "json"
    assert len(chunks) >= 1
    assert chunks[0].location_type == "json_path"
    assert "$.company" in chunks[0].location
    assert "30" in chunks[0].text


def test_html_extraction():
    html_text = """<!DOCTYPE html>
<html>
<body>
    <article id="sec-1">
        <h1>Regulatory Compliance Note</h1>
        <p>Annual audit reports must be filed within 30 days of year end.</p>
    </article>
</body>
</html>"""
    html_bytes = html_text.encode("utf-8")

    meta, chunks = parse_html(html_bytes, "portal.html")
    assert meta.filename == "portal.html"
    assert meta.file_type == "html"
    assert len(chunks) >= 1
    assert chunks[0].location_type == "element"
    assert "<" in chunks[0].location
    assert "30 days" in chunks[0].text


def test_end_to_end_docx_verification():
    import docx
    doc = docx.Document()
    doc.add_paragraph("Clause 4.1: The maximum allowable response window is 14 calendar days.")
    doc.add_paragraph("Clause 4.2: Late response incurs a fixed penalty of $5,000.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    pipeline = VerificationPipeline()
    draft = "The allowable response window is 14 days. The penalty is $50,000."
    cert = pipeline.run_verification(
        draft_text=draft,
        source_files=[("terms.docx", docx_bytes)],
        top_k=2
    )

    assert len(cert.claims) >= 2
    assert cert.sources[0].file_type == "docx"
    assert cert.sources[0].evidence_locations

    claim1 = cert.claims[0]
    assert claim1.verdict == VerdictType.SUPPORTED
    assert claim1.evidence
    assert claim1.evidence[0].location_type == "paragraph"
    assert "Paragraph" in claim1.evidence[0].location
